"""Regression tests for the medium-severity audit findings.

The HTTP layer is fake. No CouchDB server is required.
"""

import json
from pathlib import Path

import pytest

import executor as executor_module
from executor import CouchDBExecutor, ExecutorError
from parser import CountCommand, ParseError, SelectCommand, parse


class FakeResponse:
    def __init__(self, status_code, payload):
        self.status_code = status_code
        self._payload = payload
        self.text = payload if isinstance(payload, str) else json.dumps(payload)

    def json(self):
        if isinstance(self._payload, str):
            raise ValueError("not json")
        return self._payload


SHOP = [
    {"_id": "1", "_rev": "1-a", "status": "active", "n": 1, "flag": True,
     "name": "apple", "address": {"city": "Paris"}},
    {"_id": "2", "_rev": "1-b", "status": "inactive", "n": 2, "flag": False,
     "name": "app", "address": {"city": "Lyon"}},
    {"_id": "3", "_rev": "1-c", "status": None, "n": 1, "flag": 1,
     "name": "banana", "address": {"city": "Paris"}},
    {"_id": "4", "_rev": "1-d", "n": 0, "flag": 0, "name": "", "address": "nope"},
    {"_id": "5", "_rev": "1-e", "status": "active", "n": "10", "name": "apricot"},
    {"_id": "_design/idx", "_rev": "1-f", "language": "javascript"},
    {"_id": "6", "_rev": "1-g", "status": "pending", "name": "apricot2",
     "message": "error OR level = 1"},
]


class ListCouch:
    """Returns every document in one _all_docs page and records bulk posts."""

    def __init__(self, databases):
        self.databases = databases
        self.posts = []

    def request(self, method, path, **kwargs):
        if method == "GET" and path in self.databases and "/_all_docs" not in path:
            return FakeResponse(200, {"db_name": path})
        if method == "GET" and path.endswith("/_all_docs"):
            db_name = path.split("/", 1)[0]
            docs = self.databases[db_name]
            rows = [{"id": doc["_id"], "key": doc["_id"], "doc": doc} for doc in docs]
            return FakeResponse(200, {"total_rows": len(rows), "rows": rows})
        if method == "POST" and path.endswith("/_bulk_docs"):
            db_name = path.split("/", 1)[0]
            self.posts.append(kwargs["json"]["docs"])
            updated = list(self.databases[db_name])
            results = []
            for doc in kwargs["json"]["docs"]:
                updated = [existing for existing in updated if existing.get("_id") != doc.get("_id")]
                updated.append(doc)
                results.append({"ok": True, "id": doc.get("_id")})
            self.databases[db_name] = updated
            return FakeResponse(201, results)
        return FakeResponse(404, {"error": "not_found", "path": path})


def make_executor(docs=None, database="shop"):
    databases = {database: [dict(doc) for doc in (SHOP if docs is None else docs)]}
    couch = ListCouch(databases)
    executor = CouchDBExecutor()
    executor._request = couch.request
    return executor, databases, couch


def matched_ids(executor, sql):
    result = executor.execute(parse(sql))
    return [row["_id"] for row in result["rows"]]


class TestSortingAndMixedTypes:
    def test_desc_string_order_and_nulls_last(self):
        docs = list(SHOP) + [{"_id": "7", "_rev": "1-h", "name": None}]
        executor, _db, _couch = make_executor(docs)
        assert matched_ids(executor, "SELECT _id FROM shop ORDER BY name DESC;") == [
            "3", "6", "5", "1", "2", "4", "7",
        ]
        assert matched_ids(executor, "SELECT _id FROM shop ORDER BY name ASC;") == [
            "4", "2", "1", "5", "6", "3", "7",
        ]

    def test_order_by_dicts_does_not_raise(self):
        docs = [
            {"_id": "1", "_rev": "1", "meta": {"a": 1}},
            {"_id": "2", "_rev": "1", "meta": {"b": 2}},
        ]
        executor, _db, _couch = make_executor(docs)
        assert matched_ids(executor, "SELECT _id FROM shop ORDER BY meta;") == ["1", "2"]

    def test_mixed_comparison_warns_once(self):
        executor, _db, _couch = make_executor()
        assert matched_ids(executor, "SELECT _id FROM shop WHERE n > '5';") == []
        assert executor.type_warnings == ["Skipped comparisons between incompatible value types"]
        assert matched_ids(executor, "SELECT _id FROM shop WHERE flag < 2;") == ["3", "4"]
        assert executor.type_warnings == []


class TestNestedFields:
    def test_select_and_order_by_dot_path(self):
        executor, _db, _couch = make_executor()
        result = executor.execute(parse(
            "SELECT _id, address.city FROM shop ORDER BY address.city;"
        ))
        assert result["fields"] == ["_id", "address.city"]
        assert "Paris" in result["table"]
        assert "Lyon" in result["table"]
        assert [row["_id"] for row in result["rows"]] == ["2", "1", "3", "4", "5", "6"]

    def test_update_dot_path_nests_and_refuses_non_object(self):
        executor, databases, couch = make_executor()
        result = executor.execute(parse(
            "UPDATE shop SET address.city = 'Nice' WHERE _id = '1';"
        ))
        assert result["updated"] == 1
        doc = next(item for item in databases["shop"] if item["_id"] == "1")
        assert doc["address"]["city"] == "Nice"
        assert "address.city" not in doc

        with pytest.raises(ExecutorError, match="not an object"):
            executor.execute(parse(
                "UPDATE shop SET address.city = 'Nice' WHERE _id = '4';"
            ))
        assert len(couch.posts) == 1
        untouched = next(item for item in databases["shop"] if item["_id"] == "4")
        assert untouched["address"] == "nope"


class TestPagination:
    def test_pages_until_every_row_is_read(self):
        docs = [{"_id": str(i), "_rev": "1", "name": f"n{i}"} for i in range(5)]
        executor, _db, couch = make_executor(docs)
        executor._all_docs_page_size = 2

        def request(method, path, **kwargs):
            if method == "GET" and path == "shop":
                return FakeResponse(200, {})
            if method == "GET" and path.endswith("/_all_docs"):
                params = kwargs.get("params") or {}
                couch.posts.append(params)
                limit = int(params["limit"])
                start = 0
                if "startkey" in params:
                    key = json.loads(params["startkey"])
                    ids = [doc["_id"] for doc in docs]
                    start = ids.index(key) + int(params.get("skip", 0))
                page = docs[start:start + limit]
                rows = [{"id": doc["_id"], "key": doc["_id"], "doc": doc} for doc in page]
                return FakeResponse(200, {"total_rows": len(docs), "rows": rows})
            return FakeResponse(404, {})

        executor._request = request
        assert matched_ids(executor, "SELECT _id FROM shop;") == ["0", "1", "2", "3", "4"]
        assert len(couch.posts) == 3
        assert couch.posts[1]["skip"] == 1
        assert json.loads(couch.posts[1]["startkey"]) == "1"

    def test_total_rows_mismatch_aborts_before_write(self):
        executor, _db, couch = make_executor([])
        posts = []

        def request(method, path, **kwargs):
            if method == "GET" and path == "shop":
                return FakeResponse(200, {})
            if method == "GET" and path.endswith("/_all_docs"):
                return FakeResponse(200, {
                    "total_rows": 100,
                    "rows": [{
                        "id": "1", "key": "1",
                        "doc": {"_id": "1", "_rev": "1-a", "status": "active"},
                    }],
                })
            if method == "POST":
                posts.append(kwargs)
                return FakeResponse(201, [])
            return FakeResponse(404, {})

        executor._request = request
        with pytest.raises(ExecutorError, match="fetched 1 of 100"):
            executor.execute(parse("DELETE FROM shop WHERE status = 'active';"))
        with pytest.raises(ExecutorError, match="fetched 1 of 100"):
            executor.execute(parse("SELECT COUNT(*) FROM shop;"))
        assert posts == []

    def test_null_doc_counts_toward_total_rows(self):
        executor, _db, _couch = make_executor([])

        def request(method, path, **kwargs):
            return FakeResponse(200, {"total_rows": 1, "rows": [{"id": "gone", "key": "gone", "doc": None}]})

        executor._request = request
        assert executor._fetch_all_docs("shop") == []


class TestTimeoutsAndConnection:
    def test_requests_receive_a_timeout(self, monkeypatch, tmp_path):
        monkeypatch.setattr(executor_module, "_ENV_FILE", tmp_path / "missing.env")
        monkeypatch.delenv("COUCHDB_TIMEOUT", raising=False)
        captured = {}

        def fake_request(method, url, **kwargs):
            captured["timeout"] = kwargs.get("timeout")
            return FakeResponse(200, [])

        monkeypatch.setattr(executor_module.requests, "request", fake_request)
        executor = CouchDBExecutor()
        executor._request("GET", "_all_dbs")
        assert captured["timeout"] == (3.0, 60.0)

        monkeypatch.setenv("COUCHDB_TIMEOUT", "1.5,9")
        assert CouchDBExecutor()._timeout == (1.5, 9.0)
        monkeypatch.setenv("COUCHDB_TIMEOUT", "4")
        assert CouchDBExecutor()._timeout == (4.0, 4.0)
        monkeypatch.setenv("COUCHDB_TIMEOUT", "nope")
        assert CouchDBExecutor()._timeout == (3.0, 60.0)

    def test_startup_reports_401_as_auth_failure(self, monkeypatch, capsys):
        def fake_request(method, url, **kwargs):
            assert url.endswith("/_all_dbs")
            assert kwargs["timeout"] == (3.0, 60.0) or isinstance(kwargs["timeout"], tuple)
            return FakeResponse(401, {"error": "unauthorized", "reason": "Name or password is incorrect."})

        monkeypatch.delenv("COUCHDB_TIMEOUT", raising=False)
        monkeypatch.setattr(executor_module.requests, "request", fake_request)
        from main import CouchDBSQL
        CouchDBSQL()
        output = capsys.readouterr().out
        assert "HTTP 401" in output
        assert "Name or password is incorrect." in output
        assert "Could not connect" not in output

    def test_startup_reports_connection_errors(self, monkeypatch, capsys):
        def fake_request(method, url, **kwargs):
            raise executor_module.requests.ConnectionError("connection refused")

        monkeypatch.setattr(executor_module.requests, "request", fake_request)
        from main import CouchDBSQL
        CouchDBSQL()
        output = capsys.readouterr().out
        assert "Could not connect" in output
        assert "connection refused" in output


class TestExport:
    def test_sql_identifiers_are_quoted(self, tmp_path):
        weird = 'weird); DROP TABLE users;--'
        quoted = 'say"hi'
        docs = [{"_id": "1", "_rev": "1-a", weird: "x", quoted: "it's"}]
        executor, _db, _couch = make_executor(docs)
        filename = tmp_path / "out.sql"
        result = executor.execute(parse(f"EXPORT shop TO {filename};"))
        text = filename.read_text()
        assert f'"{weird}"' in text
        assert '"say""hi"' in text
        assert "CREATE TABLE \"shop\"" in text
        assert "'it''s'" in text
        assert "_rev" not in text
        assert result["fields"] == 4

    def test_excel_formula_and_nul(self, tmp_path):
        from openpyxl import load_workbook
        docs = [{
            "_id": "1", "_rev": "1-a",
            "link": '=HYPERLINK("http://evil")',
            "name": "a\x00b",
        }]
        executor, _db, _couch = make_executor(docs)
        filename = tmp_path / "out.xlsx"
        executor.execute(parse(f"EXPORT shop TO {filename};"))
        sheet = load_workbook(filename).active
        headers = [cell.value for cell in sheet[1]]
        link = sheet.cell(row=2, column=headers.index("link") + 1)
        name = sheet.cell(row=2, column=headers.index("name") + 1)
        assert link.data_type != "f"
        assert link.value.startswith("'")
        assert "\x00" not in name.value
        assert name.value == "ab"

    def test_empty_export_writes_file_and_fields(self, tmp_path, monkeypatch, capsys):
        def fake_request(method, url, **kwargs):
            if url.endswith("/_all_dbs") or url.endswith("/shop"):
                return FakeResponse(200, [])
            if url.endswith("/_all_docs"):
                return FakeResponse(200, {"total_rows": 0, "rows": []})
            return FakeResponse(404, {})

        monkeypatch.setattr(executor_module.requests, "request", fake_request)
        from main import CouchDBSQL
        shell = CouchDBSQL()
        json_path = tmp_path / "empty.json"
        sql_path = tmp_path / "empty.sql"
        xlsx_path = tmp_path / "empty.xlsx"
        shell.default(f"EXPORT shop TO {json_path};")
        shell.default(f"EXPORT shop TO {sql_path};")
        shell.default(f"EXPORT shop TO {xlsx_path};")
        output = capsys.readouterr().out
        assert "Fields: 0" in output
        assert "Error:" not in output
        assert json_path.read_text().strip() == "[]"
        assert sql_path.read_text() == "-- exported 0 documents\n"
        assert xlsx_path.is_file()


class TestCsvImport:
    def test_bom_bool_scientific_and_leading_zero(self, tmp_path):
        path = tmp_path / "products.csv"
        path.write_bytes(
            "\ufeffname,active,note,price_NUM,code_NUM,flag_BOOL\n"
            "Widget,true,true story,1e2,012,false\n"
            "Gadget,FALSE,plain,9.99,100,true\n".encode("utf-8")
        )
        executor, _db, _couch = make_executor()
        prepared = executor._prepare_csv_import(str(path), "shop")
        first, second = prepared["documents"]
        assert first["name"] == "Widget"
        assert first["active"] is True
        assert first["note"] == "true story"
        assert first["price"] == 100
        assert isinstance(first["price"], int)
        assert first["code"] == "012"
        assert first["flag"] is False
        assert second["active"] is False
        assert second["price"] == 9.99
        assert second["code"] == 100
        assert second["flag"] is True

    def test_multiline_quoted_field(self, tmp_path):
        path = tmp_path / "notes.csv"
        path.write_bytes(b'name,note\r\nAda,"line1\r\nline2"\r\n')
        executor, _db, _couch = make_executor()
        prepared = executor._prepare_csv_import(str(path), "shop")
        assert prepared["documents"] == [{"name": "Ada", "note": "line1\r\nline2"}]

    def test_non_numeric_still_rejected(self, tmp_path):
        path = tmp_path / "bad.csv"
        path.write_text("price_NUM\nN/A\n", encoding="utf-8")
        executor, _db, _couch = make_executor()
        with pytest.raises(ExecutorError, match="price_NUM"):
            executor._prepare_csv_import(str(path), "shop")


class TestImportConflicts:
    def test_conflict_reason_is_returned_without_overwriting(self, tmp_path, monkeypatch, capsys):
        posted = []
        urls = []

        def fake_request(method, url, **kwargs):
            urls.append((method, url))
            if method == "GET":
                return FakeResponse(200, [])
            posted.append(kwargs["json"]["docs"])
            return FakeResponse(201, [{
                "id": "a",
                "error": "conflict",
                "reason": "Document update conflict.",
            }])

        monkeypatch.setattr(executor_module.requests, "request", fake_request)
        monkeypatch.setattr("builtins.input", lambda prompt="": "yes")
        from main import CouchDBSQL
        shell = CouchDBSQL()
        path = tmp_path / "again.json"
        path.write_text(json.dumps([{"_id": "a", "_rev": "1-old", "name": "n"}]))
        shell.default(f"IMPORT {path} INTO shop;")
        output = capsys.readouterr().out
        assert "Errors: 1" in output
        assert "a: conflict: Document update conflict." in output
        assert "_rev" not in posted[0][0]
        assert posted[0][0]["_id"] == "a"
        assert not any(url.rstrip("/").endswith("/a") for _method, url in urls)


class TestNamesRegexAndCount:
    def test_hyphenated_names(self):
        selected = parse("SELECT user-id FROM products-v2 WHERE user-id = 'x';")
        assert isinstance(selected, SelectCommand)
        assert selected.database == "products-v2"
        assert selected.fields == ["user-id"]
        assert selected.conditions.items[0][0].field == "user-id"
        created = parse("CREATE products-v2;")
        assert created.database == "products-v2"
        updated = parse("UPDATE products-v2 SET user-id = 'a' WHERE user-id = 'b';")
        assert updated.database == "products-v2"
        assert updated.assignments == {"user-id": "a"}
        inserted = parse("INSERT INTO products-v2 (user-id, name) VALUES ('a', 'b');")
        assert inserted.rows == [{"user-id": "a", "name": "b"}]
        assert parse("CREATE MyDB;").database == "MyDB"
        with pytest.raises(ParseError):
            parse("CREATE a/b;")

        docs = [{"_id": "1", "_rev": "1", "user-id": "x"}]
        executor, _db, _couch = make_executor(docs, database="products-v2")
        result = executor.execute(parse("SELECT user-id FROM products-v2 WHERE user-id = 'x';"))
        assert result["rows"][0]["user-id"] == "x"
        assert "x" in result["table"]

    def test_invalid_regex_is_an_error(self):
        executor, _db, _couch = make_executor()
        with pytest.raises(ExecutorError, match="Invalid regular expression"):
            executor.execute(parse("SELECT _id FROM shop WHERE name ~ '[';"))
        assert matched_ids(executor, "SELECT _id FROM shop WHERE name ~ 'app';") == ["1", "2"]

    def test_count_rejects_limit(self):
        with pytest.raises(ParseError, match="LIMIT"):
            parse("SELECT COUNT(*) FROM products LIMIT 5;")
        with pytest.raises(ParseError, match="ORDER BY"):
            parse("SELECT COUNT(*) FROM products ORDER BY name;")
        counted = parse("SELECT COUNT(*) FROM products;")
        assert isinstance(counted, CountCommand)
        assert not isinstance(counted, SelectCommand)
        distinct = parse("SELECT COUNT(DISTINCT address.city) FROM products;")
        assert distinct.distinct_field == "address.city"


class TestEnvFile:
    def test_default_path_is_beside_the_module(self):
        assert executor_module._ENV_FILE == Path(executor_module.__file__).resolve().parent / ".env"

    def test_loads_explicit_env_file_not_the_working_directory(self, monkeypatch, tmp_path):
        env_file = tmp_path / "proj.env"
        env_file.write_text(
            "COUCHDB_URL=http://from-file:5984\nCOUCHDB_USER=ada\nCOUCHDB_PASS=secret\n"
        )
        empty = tmp_path / "empty"
        empty.mkdir()
        (empty / ".env").write_text(
            "COUCHDB_URL=http://from-cwd:5984\nCOUCHDB_USER=cwd\nCOUCHDB_PASS=nope\n"
        )
        monkeypatch.chdir(empty)
        monkeypatch.setattr(executor_module, "_ENV_FILE", env_file)
        for name in ("COUCHDB_URL", "COUCHDB_USER", "COUCHDB_PASS", "COUCHDB_TIMEOUT"):
            monkeypatch.delenv(name, raising=False)
        executor = CouchDBExecutor()
        assert executor.url == "http://from-file:5984"
        assert executor.user == "ada"


def test_license_gitignore_and_ci_exist():
    root = Path(__file__).resolve().parents[1]
    assert "MIT License" in (root / "LICENSE").read_text(encoding="utf-8")
    assert (root / ".gitignore").is_file()
    workflow = (root / ".github" / "workflows" / "tests.yml").read_text(encoding="utf-8")
    assert "pytest" in workflow
